"""Generate .xlsx attendance report."""
from io import BytesIO
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autofit(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)


def generate_excel(
    user_id: int,
    username: str,
    all_attendance: List[Dict[str, Any]],
    stats_raw: List[Dict[str, Any]],
) -> bytes:
    """Return .xlsx bytes.

    *stats_raw*  – list from storage.get_attendance_stats()
    *all_attendance* – list from storage.get_all_attendance_log()
    """
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------------
    # Sheet 1 – Summary per discipline
    # ------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Sumar"

    # Aggregate stats_raw into {subject: {total, prezente, absente, types: {C/L/P: ...}}}
    aggregated: Dict[str, Any] = {}
    for row in stats_raw:
        subj = row["subject"]
        st = row["session_type"]
        if subj not in aggregated:
            aggregated[subj] = {"total": 0, "prezente": 0, "absente": 0, "types": {}}
        aggregated[subj]["total"] += row["total"]
        aggregated[subj]["prezente"] += row["prezente"]
        aggregated[subj]["absente"] += row["absente"]
        aggregated[subj]["types"][st] = row

    headers1 = ["Disciplina", "Total", "Prezențe", "Absențe", "% Prezență", "Detalii C", "Detalii L", "Detalii P"]
    for col, h in enumerate(headers1, 1):
        ws1.cell(row=1, column=col, value=h)
    _style_header(ws1, 1, len(headers1))
    ws1.row_dimensions[1].height = 20

    for r, (subj, data) in enumerate(sorted(aggregated.items()), start=2):
        total = data["total"]
        prez = data["prezente"]
        pct = (prez / total * 100) if total else 0.0
        ws1.cell(row=r, column=1, value=subj)
        ws1.cell(row=r, column=2, value=total)
        ws1.cell(row=r, column=3, value=prez)
        ws1.cell(row=r, column=4, value=data["absente"])
        ws1.cell(row=r, column=5, value=f"{pct:.1f}%")
        for col_offset, stype in enumerate(("C", "L", "P"), start=6):
            t = data["types"].get(stype)
            if t:
                ws1.cell(row=r, column=col_offset,
                         value=f"{t['prezente']}/{t['total']}")

    _autofit(ws1)

    # ------------------------------------------------------------------
    # Sheet 2 – Full attendance log
    # ------------------------------------------------------------------
    ws2 = wb.create_sheet("Log complet")
    headers2 = [
        "Data", "Zi", "Ora start", "Ora end",
        "Disciplina", "Tip", "Grupă", "Sală",
        "Status", "Utilizator", "Înregistrat la",
    ]
    for col, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=col, value=h)
    _style_header(ws2, 1, len(headers2))
    ws2.row_dimensions[1].height = 20

    for r, rec in enumerate(all_attendance, start=2):
        ws2.cell(row=r, column=1,  value=rec.get("date"))
        ws2.cell(row=r, column=2,  value=rec.get("day"))
        ws2.cell(row=r, column=3,  value=rec.get("time_start"))
        ws2.cell(row=r, column=4,  value=rec.get("time_end"))
        ws2.cell(row=r, column=5,  value=rec.get("subject"))
        ws2.cell(row=r, column=6,  value=rec.get("session_type"))
        ws2.cell(row=r, column=7,  value=rec.get("group_code"))
        ws2.cell(row=r, column=8,  value=rec.get("room"))
        status = rec.get("status", "")
        ws2.cell(row=r, column=9,  value=status)
        ws2.cell(row=r, column=10, value=username)
        ws2.cell(row=r, column=11, value=rec.get("recorded_at"))

        # Color rows by status
        fill_color = "C6EFCE" if status == "prezent" else "FFC7CE" if status == "absent" else None
        if fill_color:
            fill = PatternFill("solid", fgColor=fill_color)
            for col in range(1, len(headers2) + 1):
                ws2.cell(row=r, column=col).fill = fill

    _autofit(ws2)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
